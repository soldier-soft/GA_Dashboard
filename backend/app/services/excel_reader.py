import os
import openpyxl
from datetime import datetime
from typing import Dict, Any, List, Tuple
from app.models import TableData, TableRow, DiscrepancyNotice

DEFAULT_EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "data",
    "Station_Assessment_Dashboard.xlsx"
)

# Fallback path if not found in data/
FALLBACK_ORIGINAL_PATH = r"C:\Users\Elango\OneDrive\Apps\Documents\Station_Assessment_Dashboard_Exact_Layout (1).xlsx"

class ExcelReaderService:
    def __init__(self, excel_path: str = None):
        self.excel_path = excel_path or (
            DEFAULT_EXCEL_PATH if os.path.exists(DEFAULT_EXCEL_PATH) else FALLBACK_ORIGINAL_PATH
        )
        self.cached_tables: Dict[str, TableData] = {}
        self.cached_discrepancies: List[DiscrepancyNotice] = []
        self.last_read_time: datetime = datetime.now()
        self.file_mtime: float = 0

    def get_file_info(self) -> Dict[str, Any]:
        if not os.path.exists(self.excel_path):
            return {
                "exists": False,
                "path": self.excel_path,
                "error": "Workbook file not found on disk"
            }
        mtime = os.path.getmtime(self.excel_path)
        mtime_str = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S")
        return {
            "exists": True,
            "path": self.excel_path,
            "file_name": os.path.basename(self.excel_path),
            "size_bytes": os.path.getsize(self.excel_path),
            "modified_time": mtime_str,
            "raw_mtime": mtime
        }

    def read_workbook(self, force: bool = False) -> Tuple[Dict[str, TableData], List[DiscrepancyNotice]]:
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel workbook not found at: {self.excel_path}")

        current_mtime = os.path.getmtime(self.excel_path)
        if not force and self.cached_tables and current_mtime == self.file_mtime:
            return self.cached_tables, self.cached_discrepancies

        # Load with both formulas and evaluated values
        wb_form = openpyxl.load_workbook(self.excel_path, data_only=False)
        wb_val = openpyxl.load_workbook(self.excel_path, data_only=True)

        if "Data" not in wb_val.sheetnames:
            raise ValueError(f"Required sheet 'Data' not found in {wb_val.sheetnames}")

        ws_data_form = wb_form["Data"]
        ws_data_val = wb_val["Data"]

        tables: Dict[str, TableData] = {}
        discrepancies: List[DiscrepancyNotice] = []

        # -------------------------------------------------------------
        # Table 1 — Assessment Completion
        # Data range: Data!A4:C7
        # -------------------------------------------------------------
        t1_rows: List[TableRow] = []
        total_impacted = 0.0
        completed_val = 0.0
        inprogress_val = 0.0

        for r in range(5, 8):
            cat = str(ws_data_val.cell(r, 1).value or "").strip()
            val = float(ws_data_val.cell(r, 2).value or 0.0)
            form_pct = ws_data_form.cell(r, 3).value
            val_pct = ws_data_val.cell(r, 3).value

            if r == 5:
                total_impacted = val
                row_id = "t1-impacted"
            elif r == 6:
                completed_val = val
                row_id = "t1-completed"
            else:
                inprogress_val = val
                row_id = "t1-inprogress"

            # Compute percentage safely from baseline (B5)
            calc_pct = (val / total_impacted) if total_impacted > 0 else 0.0
            display_pct = float(val_pct) if isinstance(val_pct, (int, float)) else calc_pct

            t1_rows.append(TableRow(
                id=row_id,
                category=cat,
                value=val,
                percentage=round(display_pct, 4),
                excel_cell_ref=f"Data!A{r}:C{r}",
                excel_formula=str(form_pct) if form_pct else None,
                is_placeholder=False
            ))

        # Check for discrepancy in Table 1
        sum_components = completed_val + inprogress_val
        if sum_components > total_impacted and total_impacted > 0:
            diff = int(sum_components - total_impacted)
            discrepancies.append(DiscrepancyNotice(
                id="disc-t1-sum-overflow",
                severity="warning",
                title="Assessment Counts Exceed Total Impacted Stations",
                description=(
                    f"Completed Stations ({int(completed_val)}) + In Progress Stations ({int(inprogress_val)}) = "
                    f"{int(sum_components)}, exceeding Total LHD Impacted Stations ({int(total_impacted)}) by {diff}. "
                    "Displaying faithful Excel model values without forcing reconciliation."
                ),
                table_id="table-1",
                raw_details={
                    "total_impacted": total_impacted,
                    "completed": completed_val,
                    "in_progress": inprogress_val,
                    "excess": diff
                }
            ))

        tables["table-1"] = TableData(
            id="table-1",
            title="Table 1 — Assessment Completion",
            subtitle="LHD Impacted Station Progress & Status",
            category_header="Assessment Status",
            value_header="Stations",
            is_excel_source=True,
            is_editable=False,
            chart_type="donut",
            chart_color="#2563eb",
            description="Authoritative station assessment counts and completion rate from Excel Data sheet.",
            rows=t1_rows,
            total=total_impacted,
            summary={
                "impacted": total_impacted,
                "completed": completed_val,
                "in_progress": inprogress_val,
                "completion_rate": round((completed_val / total_impacted) * 100, 2) if total_impacted > 0 else 0.0
            }
        )

        # -------------------------------------------------------------
        # Table 2 — WPI (Workplace Integration)
        # Data range: Data!A10:B14
        # -------------------------------------------------------------
        t2_rows: List[TableRow] = []
        t2_total = 0.0
        for r in range(11, 15):
            cat = str(ws_data_val.cell(r, 1).value or "").strip()
            val = float(ws_data_val.cell(r, 2).value or 0.0)
            t2_total += val
            t2_rows.append(TableRow(
                id=f"t2-row-{r-10}",
                category=cat,
                value=val,
                percentage=0.0, # calculated below
                excel_cell_ref=f"Data!A{r}:B{r}",
                is_placeholder=False
            ))

        # populate percentage share
        for row in t2_rows:
            row.percentage = round((row.value / t2_total), 4) if t2_total > 0 else 0.0

        tables["table-2"] = TableData(
            id="table-2",
            title="Table 2 — WPI",
            subtitle="Workplace Integration by Category",
            category_header="WPI Category",
            value_header="Data",
            is_excel_source=True,
            is_editable=False,
            chart_type="column",
            chart_color="#0284c7",
            description="Assessment distribution across Health & Safety, Customer Experience, Logistics, and Workplace Organization.",
            rows=t2_rows,
            total=t2_total,
            summary={
                "highest_category": max(t2_rows, key=lambda x: x.value).category if t2_rows else "",
                "total_wpi_points": t2_total
            }
        )

        # -------------------------------------------------------------
        # Table 3 — Health & Safety Status
        # Data range: Data!E4:F7
        # -------------------------------------------------------------
        t3_rows: List[TableRow] = []
        t3_total_item = 0.0
        t3_inprog = 0.0
        t3_closed = 0.0

        for r in range(5, 8):
            cat = str(ws_data_val.cell(r, 5).value or "").strip()
            val = float(ws_data_val.cell(r, 6).value or 0.0)
            if cat.lower() == "total":
                t3_total_item = val
                row_id = "t3-total"
            elif "progress" in cat.lower():
                t3_inprog = val
                row_id = "t3-inprog"
            else:
                t3_closed = val
                row_id = "t3-closed"

            t3_rows.append(TableRow(
                id=row_id,
                category=cat,
                value=val,
                percentage=0.0,
                excel_cell_ref=f"Data!E{r}:F{r}",
                is_placeholder=False
            ))

        # Check reconciliation: In progress + Closed == Total
        if t3_inprog + t3_closed != t3_total_item and t3_total_item > 0:
            discrepancies.append(DiscrepancyNotice(
                id="disc-t3-sum-mismatch",
                severity="info",
                title="Health & Safety Sub-status Count Variance",
                description=f"In progress ({t3_inprog}) + Closed ({t3_closed}) != Total ({t3_total_item}).",
                table_id="table-3",
                raw_details={"total": t3_total_item, "in_progress": t3_inprog, "closed": t3_closed}
            ))

        # Calculate percentage of total for status items
        for row in t3_rows:
            if t3_total_item > 0:
                row.percentage = round(row.value / t3_total_item, 4)

        tables["table-3"] = TableData(
            id="table-3",
            title="Table 3 — Health & Safety",
            subtitle="Safety Assessment Status Breakdown",
            category_header="Health and Safety",
            value_header="Data",
            is_excel_source=True,
            is_editable=False,
            chart_type="column",
            chart_color="#10b981",
            description="Resolution status for shop-floor Health and Safety assessments.",
            rows=t3_rows,
            total=t3_total_item,
            summary={
                "closure_rate": round((t3_closed / t3_total_item) * 100, 2) if t3_total_item > 0 else 0.0,
                "in_progress": t3_inprog,
                "closed": t3_closed
            }
        )

        # -------------------------------------------------------------
        # Table 4 — Action Category
        # Data range: Data!H4:I7 (Facility, Process, Training)
        # -------------------------------------------------------------
        t4_rows: List[TableRow] = []
        t4_total = 0.0
        for r in range(5, 8):
            cat = str(ws_data_val.cell(r, 8).value or "").strip()
            val = float(ws_data_val.cell(r, 9).value or 0.0)
            t4_total += val
            t4_rows.append(TableRow(
                id=f"t4-row-{r-4}",
                category=cat,
                value=val,
                percentage=0.0,
                excel_cell_ref=f"Data!H{r}:I{r}",
                is_placeholder=False
            ))

        for row in t4_rows:
            row.percentage = round((row.value / t4_total), 4) if t4_total > 0 else 0.0

        tables["table-4"] = TableData(
            id="table-4",
            title="Table 4 — Action Category",
            subtitle="Corrective Actions by Functional Discipline",
            category_header="Action Category",
            value_header="Data",
            is_excel_source=True,
            is_editable=False,
            chart_type="bar",
            chart_color="#f59e0b",
            description="Categorization of corrective actions across Facility, Process, and Training (Template standard for Tables 5-7).",
            rows=t4_rows,
            total=t4_total,
            summary={
                "dominant_category": max(t4_rows, key=lambda x: x.value).category if t4_rows else "",
                "total_actions": t4_total
            }
        )

        self.cached_tables = tables
        self.cached_discrepancies = discrepancies
        self.last_read_time = datetime.now()
        self.file_mtime = current_mtime

        return tables, discrepancies
